from __future__ import annotations

from io import BytesIO

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import HTMLResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.services.project_persistence import (
    get_project,
    get_search_run,
    list_projects,
    list_search_runs,
)
from app.worker import celery_app


router = APIRouter(prefix="/projects", tags=["projects"])


def _stage_label(value):
    return {"A":"A – Ранняя стадия", "B":"B – Подготовка площадки", "V":"V – Участок определён", "G":"G – Строительство началось", "unknown":"Стадия уточняется"}.get(value, "Стадия уточняется")

def _land_label(value):
    return {"confirmed_needed":"Требуется земельный участок", "high_probability":"Вероятно требуется участок", "land_defined":"Участок уже определён", "unknown":"Статус участка уточняется"}.get(value, "Статус участка уточняется")

def _priority_label(value):
    return {"A_hot":"Высокий", "A_priority":"Высокий", "B_work":"Перспективный", "B_strong":"Перспективный", "C_verify":"Требует проверки", "D_research":"Низкий", "D_low":"Низкий"}.get(value, "Не определён")

def _action_label(value):
    return {"resolve_company":"Установить компанию-инвестора", "contact_now":"Связаться с ЛПР", "find_decision_maker":"Найти ЛПР и контакты", "verify_project_status":"Проверить актуальный статус проекта", "research_later":"Повторно проверить позже", "deep_verify_now":"Срочно проверить проект", "verify_and_enrich":"Проверить и собрать контакты", "verify_if_capacity":"Проверить статус", "keep_low_priority":"Повторно проверить позже", "manual_verification":"Проверить вручную"}.get(value, "Проверить лид")


@router.get("")
def projects_list(
    bucket: str | None = None,
    min_score: int | None = Query(default=None, ge=0, le=100),
    min_sales_score: int | None = Query(default=None, ge=0, le=100),
    sales_priority: str | None = None,
    stage: str | None = None,
    company: str | None = None,
    project_type: str | None = None,
    limit: int = Query(default=100, ge=1, le=500),
    offset: int = Query(default=0, ge=0),
):
    projects = list_projects(
        bucket=bucket,
        min_score=min_score,
        min_sales_score=min_sales_score,
        sales_priority=sales_priority,
        stage=stage,
        company=company,
        project_type=project_type,
        limit=limit,
        offset=offset,
    )
    return {"count": len(projects), "limit": limit, "offset": offset, "projects": projects}


@router.get("/search-runs")
def search_runs(limit: int = Query(default=50, ge=1, le=200)):
    runs = list_search_runs(limit=limit)
    return {"count": len(runs), "runs": runs}


@router.post("/search-runs/start")
def start_search_run():
    task = celery_app.send_task("qualification.discovery_test")
    return {
        "status": "queued",
        "task_id": task.id,
        "task_status_url": f"/projects/tasks/{task.id}",
    }


@router.get("/search-runs/{run_id}")
def search_run_detail(run_id: str):
    run = get_search_run(run_id)
    if not run:
        raise HTTPException(status_code=404, detail="Search run not found")
    return run


@router.get("/tasks/{task_id}")
def task_status(task_id: str):
    result = celery_app.AsyncResult(task_id)
    payload = {"task_id": task_id, "status": result.status}
    if result.successful():
        payload["result"] = result.result
    elif result.failed():
        payload["error"] = str(result.result)
    return payload


@router.get("/export.xlsx")
def projects_export_xlsx(
    bucket: str | None = None,
    min_score: int | None = Query(default=None, ge=0, le=100),
    min_sales_score: int | None = Query(default=None, ge=0, le=100),
    stage: str | None = None,
    company: str | None = None,
    project_type: str | None = None,
):
    projects = list_projects(
        bucket=bucket,
        min_score=min_score,
        min_sales_score=min_sales_score,
        stage=stage,
        company=company,
        project_type=project_type,
        limit=500,
        offset=0,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Лиды"
    headers = [
        "ID", "Статус лида", "Sales Score", "Project Score", "Приоритет продажи",
        "Компания", "Юр. лицо", "ИНН", "Сайт", "Тип проекта", "Описание",
        "Локация", "Инвестиции, ₽", "Стадия проекта", "Земельный участок", "Статус проекта",
        "Уверенность данных", "Что делать", "Первое обнаружение", "Последнее обнаружение",
    ]
    ws.append(headers)
    for p in projects:
        ws.append([
            p.get("id"), p.get("bucket"), p.get("sales_score"), p.get("project_score") or p.get("lead_score"),
            _priority_label(p.get("sales_priority") or p.get("priority")), p.get("company_name"), p.get("legal_name"), p.get("inn"), p.get("website"),
            p.get("project_type"), p.get("project_summary"), p.get("location"), p.get("investment_rub"),
            _stage_label(p.get("stage")), _land_label(p.get("land_status")), p.get("signal_status"), p.get("confidence"),
            _action_label(p.get("recommended_action")), p.get("first_seen_at"), p.get("last_seen_at"),
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [38,18,12,12,16,28,34,14,30,38,70,42,18,10,18,18,14,24,22,22]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="zemlya-leads.xlsx"'},
    )


DASHBOARD_HTML = r'''<!doctype html>
<html lang="ru">
<head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Zemlya AI – лиды</title>
<style>
:root{--bg:#f5f7fa;--card:#fff;--text:#18212f;--muted:#6b7280;--line:#e5e7eb;--accent:#173b67;--ok:#0f766e;--warn:#b45309}
*{box-sizing:border-box}body{margin:0;font:14px Inter,Arial,sans-serif;background:var(--bg);color:var(--text)}
header{background:#fff;border-bottom:1px solid var(--line);padding:18px 24px;display:flex;align-items:center;justify-content:space-between;position:sticky;top:0;z-index:4}
h1{font-size:20px;margin:0}.sub{color:var(--muted);font-size:12px;margin-top:3px}.wrap{padding:20px 24px;max-width:1700px;margin:auto}
.toolbar,.stats{display:flex;gap:10px;flex-wrap:wrap;margin-bottom:14px}.card{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:12px 14px}
.stat{min-width:130px}.stat b{font-size:22px;display:block;margin-top:3px}input,select,button,a.btn{border:1px solid #cfd5dd;border-radius:8px;padding:9px 10px;background:#fff;color:var(--text);text-decoration:none}
button.primary{background:var(--accent);color:#fff;border-color:var(--accent);cursor:pointer}button:disabled{opacity:.55}.grow{flex:1;min-width:220px}
.tablebox{background:#fff;border:1px solid var(--line);border-radius:12px;overflow:auto}table{width:100%;border-collapse:collapse;min-width:1350px}th,td{padding:10px;border-bottom:1px solid var(--line);text-align:left;vertical-align:top}th{background:#fafbfc;font-size:12px;position:sticky;top:0}tr:hover{background:#fafcff}.score{font-weight:700;font-size:16px}.muted{color:var(--muted)}
.badge{display:inline-block;padding:3px 7px;border-radius:999px;background:#edf2f7;font-size:11px;white-space:nowrap}.hot{background:#d1fae5;color:#065f46}.verify{background:#fef3c7;color:#92400e}
.summary{max-width:380px}.company{font-weight:650}.click{cursor:pointer}.modal{display:none;position:fixed;inset:0;background:#0006;z-index:20;padding:5vh 8vw}.modal.open{display:block}.panel{background:#fff;border-radius:14px;max-width:1100px;margin:auto;max-height:90vh;overflow:auto;padding:22px}.panel h2{margin-top:0}.grid{display:grid;grid-template-columns:1fr 1fr;gap:12px}.wide{grid-column:1/-1}.contact{padding:9px;border:1px solid var(--line);border-radius:8px;margin:6px 0}.status{font-size:12px;color:var(--muted)}
@media(max-width:800px){.wrap{padding:12px}.grid{grid-template-columns:1fr}.modal{padding:2vh 2vw}}
</style></head>
<body>
<header><div><h1>Zemlya AI – потенциальные клиенты</h1><div class="sub">Актуальные инвестиционные проекты Московской области – идти по списку сверху вниз</div></div><div class="status" id="runStatus">Готово</div></header>
<div class="wrap">
<div class="stats"><div class="card stat">Актуальных проектов<b id="sAll">–</b></div><div class="card stat">Готовы к контакту<b id="sReady">–</b></div><div class="card stat">Высокий потенциал<b id="sHot">–</b></div><div class="card stat">Нужна проверка<b id="sVerify">–</b></div></div>
<details class="card" style="margin-bottom:14px"><summary style="cursor:pointer;font-weight:700">Как читать список?</summary><div style="margin-top:10px;line-height:1.55"><b>A – Ранняя стадия:</b> проект объявлен, участок требуется или ещё не подтверждён. <b>B – Подготовка площадки:</b> территория понятна, но участок выбирают или оформляют. <b>V – Участок определён:</b> земля уже выбрана или предоставлена. <b>G – Строительство началось:</b> обычно низкий приоритет для земельной услуги.<br><br><b>Project Score</b> – насколько проект соответствует услугам «Земля без торгов». <b>Sales Score</b> – насколько реально можно начинать продажу сейчас: <b>90–100 – найден ЛПР с прямым контактом</b>, <b>70–89 – сильный лид, но контакт нужно дособрать</b>, <b>50–69 – сначала проверить данные</b>, <b>0–49 – низкий приоритет</b>.</div></details>
<div class="card" style="margin-bottom:14px;line-height:1.55"><b>Как пользоваться:</b> идите по списку сверху вниз. Сначала звоните/пишите лидам 90–100. Для 70–89 сначала откройте карточку и дособерите прямой контакт ЛПР. 50–69 – перепроверьте статус проекта. Ниже 50 – только при наличии свободного времени.</div>
<div class="toolbar card">
<select id="priority"><option value="">Любой приоритет</option><option value="A_hot">Готов к контакту</option><option value="B_work">Высокий потенциал</option><option value="C_verify">Требует проверки</option><option value="D_research">Низкий приоритет</option></select>
<select id="stage"><option value="">Любая стадия</option><option value="A">A – Ранняя стадия</option><option value="B">B – Подготовка площадки</option><option value="V">V – Участок определён</option><option value="G">G – Строительство началось</option><option value="unknown">Стадия уточняется</option></select>
<input id="company" class="grow" placeholder="Поиск по компании">
<input id="minSales" type="number" min="0" max="100" placeholder="Sales Score от 0 до 100">
<button onclick="load()">Применить фильтры</button><button class="primary" id="runBtn" onclick="startRun()">Запустить новый поиск</button>
<a class="btn" href="/projects/export.xlsx">Скачать Excel</a>
</div>
<div class="tablebox"><table><thead><tr><th>Sales Score</th><th>Project Score</th><th>Приоритет</th><th>Компания</th><th>Проект</th><th>Локация</th><th>Инвестиции</th><th>Стадия проекта</th><th>Земельный участок</th><th>Что делать</th></tr></thead><tbody id="rows"></tbody></table></div>
</div>
<div class="modal" id="modal" onclick="if(event.target===this)closeModal()"><div class="panel"><button style="float:right" onclick="closeModal()">Закрыть</button><div id="detail">Загрузка…</div></div></div>
<script>
const esc=s=>String(s??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
const money=v=>v==null?'–':new Intl.NumberFormat('ru-RU',{maximumFractionDigits:0}).format(v)+' ₽';
const stageLabel=v=>({A:'A – Ранняя стадия',B:'B – Подготовка площадки',V:'V – Участок определён',G:'G – Строительство началось',unknown:'Стадия уточняется'}[v]||'Стадия уточняется');
const landLabel=v=>({confirmed_needed:'Требуется земельный участок',high_probability:'Вероятно требуется участок',land_defined:'Участок уже определён',unknown:'Статус участка уточняется'}[v]||'Статус участка уточняется');
const priorityLabel=v=>({A_hot:'Готов к контакту',A_priority:'Готов к контакту',B_work:'Высокий потенциал',B_strong:'Высокий потенциал',C_verify:'Требует проверки',D_research:'Низкий приоритет',D_low:'Низкий приоритет'}[v]||'Не определён');
const companyConfidenceLabel=v=>({high:'Подтверждено',medium:'Вероятно подтверждено',low:'Слабое подтверждение',unresolved:'Не подтверждено'}[v]||'Не подтверждено');
const actionLabel=v=>({resolve_company:'Установить компанию-инвестора',contact_now:'Связаться с ЛПР',find_decision_maker:'Найти ЛПР и контакты',verify_project_status:'Проверить актуальный статус проекта',research_later:'Повторно проверить позже',deep_verify_now:'Срочно проверить проект',verify_and_enrich:'Проверить и собрать контакты',verify_if_capacity:'Проверить статус',keep_low_priority:'Повторно проверить позже',manual_verification:'Проверить вручную'}[v]||'Проверить лид');
async function load(){let q=new URLSearchParams();for(let [id,key] of [['priority','sales_priority'],['stage','stage'],['company','company'],['minSales','min_sales_score']]){let v=document.getElementById(id).value;if(v)q.set(key,v)}q.set('limit','200');let d=await fetch('/projects?'+q).then(r=>r.json());render(d.projects||[])}
function render(ps){document.getElementById('sAll').textContent=ps.length;document.getElementById('sReady').textContent=ps.filter(x=>(x.sales_score??0)>=90).length;document.getElementById('sHot').textContent=ps.filter(x=>(x.sales_score??0)>=70&&(x.sales_score??0)<90).length;document.getElementById('sVerify').textContent=ps.filter(x=>(x.sales_score??0)>=50&&(x.sales_score??0)<70||x.bucket==='verification_pool').length;document.getElementById('rows').innerHTML=ps.map(p=>`<tr class="click" onclick="openProject('${p.id}')"><td class="score">${p.sales_score??'–'}</td><td>${p.project_score??p.lead_score??'–'}</td><td><span class="badge ${p.sales_priority==='A_hot'?'hot':p.bucket==='verification_pool'?'verify':''}">${priorityLabel(p.sales_priority||p.priority)}</span></td><td><div class="company">${esc(p.company_name||'Компания не установлена')}</div><div class="muted">${esc(p.inn?'ИНН '+p.inn:'')}</div></td><td class="summary"><b>${esc(p.project_type||'–')}</b><div class="muted">${esc(p.project_summary||'')}</div></td><td>${esc(p.location||'–')}</td><td>${money(p.investment_rub)}</td><td>${stageLabel(p.stage)}</td><td>${landLabel(p.land_status)}</td><td>${actionLabel(p.recommended_action)}</td></tr>`).join('')}
async function openProject(id){document.getElementById('modal').classList.add('open');let p=await fetch('/projects/'+id).then(r=>r.json());let contacts=(p.contacts||[]).map(c=>`<div class="contact"><b>${esc(c.name||'Корпоративный контакт')}</b> – ${esc(c.role||'')}<br>${esc(c.email||'')} ${esc(c.phone||'')}<br><a href="${esc(c.source_url)}" target="_blank">источник</a></div>`).join('')||'<span class="muted">Подтверждённые контакты пока не найдены</span>';let sources=(p.sources||[]).map(s=>`<div><a href="${esc(s.url)}" target="_blank">${esc(s.title||s.domain||s.url)}</a></div>`).join('');document.getElementById('detail').innerHTML=`<h2>${esc(p.company_name||p.project_type||'Проект')}</h2><div class="grid"><div class="card"><b>Sales Score – готовность лида к продаже</b><div class="score">${p.sales_score??'–'} / 100</div></div><div class="card"><b>Project Score – соответствие нашим услугам</b><div class="score">${p.project_score??p.lead_score??'–'} / 100</div></div><div class="card wide"><b>Что происходит</b><p>${esc(p.project_summary||'–')}</p><div>${esc(p.location||'–')} – ${money(p.investment_rub)}</div></div><div class="card"><b>Компания</b><p>${esc(p.company_name||'Не установлена')}</p><div>${esc(p.legal_name||'')}</div><div>${esc(p.inn?'ИНН '+p.inn:'')} ${esc(p.ogrn?'ОГРН '+p.ogrn:'')}</div>${p.website?`<a href="${esc(p.website)}" target="_blank">Сайт</a>`:''}</div><div class="card"><b>Что делать</b><p>${priorityLabel(p.sales_priority||p.priority)}</p><div>Стадия: ${stageLabel(p.stage)}</div><div>Земля: ${landLabel(p.land_status)}</div><div>Следующий шаг: ${actionLabel(p.recommended_action)}</div></div><div class="card wide"><b>ЛПР и контакты</b>${contacts}</div><div class="card wide"><b>Источники</b>${sources}</div></div>`}
function closeModal(){document.getElementById('modal').classList.remove('open')}
async function startRun(){let b=document.getElementById('runBtn');b.disabled=true;document.getElementById('runStatus').textContent='Идёт поиск новых лидов…';try{let x=await fetch('/projects/search-runs/start',{method:'POST'}).then(r=>r.json());poll(x.task_id)}catch(e){document.getElementById('runStatus').textContent='Не удалось запустить поиск';b.disabled=false}}
async function poll(id){let x=await fetch('/projects/tasks/'+id).then(r=>r.json());document.getElementById('runStatus').textContent='Поиск: '+x.status;if(['SUCCESS','FAILURE'].includes(x.status)){document.getElementById('runBtn').disabled=false;if(x.status==='SUCCESS')load();return}setTimeout(()=>poll(id),4000)}
load();
</script></body></html>'''


@router.get("/dashboard", response_class=HTMLResponse)
def dashboard():
    return DASHBOARD_HTML


@router.get("/{project_id}")
def project_detail(project_id: str):
    project = get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="Project not found")
    return project
